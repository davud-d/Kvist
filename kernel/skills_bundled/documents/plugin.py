# -*- coding: utf-8 -*-
"""Kvist documents skill: PDF pages, spreadsheet ranges, CSV rows and DOCX paragraphs with references."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

DEFAULT_CHARS = 40_000
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_RANGE_RE = re.compile(r"^\s*(\d+)\s*(?:-\s*(\d+))?\s*$")
_A1_RE = re.compile(r"^([A-Z]{1,3})(\d+)(?::([A-Z]{1,3})(\d+))?$")


def describe() -> dict[str, Any]:
    return {"tools": ["info", "read"]}


def execute(tool_name: str, arguments: dict[str, Any], context: Any) -> dict[str, Any]:
    source = Path(context.files.read_bytes(arguments["path"]))
    kind = kind_of(arguments["path"])
    if tool_name == "info":
        return {"path": arguments["path"], "type": kind, **INFO[kind](source)}
    if tool_name == "read":
        limit = int(arguments.get("max_chars") or DEFAULT_CHARS)
        result = READ[kind](source, arguments)
        return {"path": arguments["path"], "type": kind, **bounded(result, limit), "untrusted": True}
    raise ValueError(f"unknown tool {tool_name}")


def kind_of(path: str) -> str:
    suffix = Path(path).suffix.lower()
    kinds = {".pdf": "pdf", ".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "csv", ".tsv": "csv", ".docx": "docx"}
    if suffix not in kinds:
        raise ValueError(f"{suffix or 'this file type'} is not supported; pdf, xlsx, csv/tsv and docx are")
    return kinds[suffix]


def page_numbers(spec: str, total: int) -> list[int]:
    """1-based page numbers from "1-3,7"; empty spec means every page."""
    if not spec:
        return list(range(1, total + 1))
    pages: list[int] = []
    for part in spec.split(","):
        match = _RANGE_RE.match(part)
        if not match:
            raise ValueError(f"{part!r} is not a page or a range like 2-5")
        first, last = int(match.group(1)), int(match.group(2) or match.group(1))
        if first < 1 or last < first:
            raise ValueError(f"{part!r} is not a valid range")
        pages.extend(number for number in range(first, last + 1) if number <= total)
    return list(dict.fromkeys(pages))


def bounded(result: dict[str, Any], limit: int) -> dict[str, Any]:
    used = 0
    kept = []
    for block in result["blocks"]:
        size = len(str(block))
        if used + size > limit and kept:
            result["truncated"] = True
            result["next"] = block.get("ref")
            break
        kept.append(block)
        used += size
    result["blocks"] = kept
    result.setdefault("truncated", False)
    return result


# --- PDF -------------------------------------------------------------------


def pdf_info(source: Path) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(str(source))
    meta = reader.metadata or {}
    return {"pages": len(reader.pages), "encrypted": reader.is_encrypted,
            "title": str(meta.get("/Title") or ""), "author": str(meta.get("/Author") or "")}


def pdf_read(source: Path, arguments: dict[str, Any]) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(str(source))
    blocks = []
    empty = 0
    for number in page_numbers(str(arguments.get("pages") or ""), len(reader.pages)):
        text = reader.pages[number - 1].extract_text() or ""
        empty += not text.strip()
        blocks.append({"ref": f"page {number}", "page": number, "text": text})
    result: dict[str, Any] = {"pages_total": len(reader.pages), "blocks": blocks}
    if blocks and empty == len(blocks):
        result["note"] = "These pages have no text layer (a scan?); OCR is not part of this skill."
    return result


# --- spreadsheets ------------------------------------------------------------


def xlsx_info(source: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    book = load_workbook(str(source), read_only=True, data_only=True)
    sheets = [{"name": sheet.title, "range": sheet.calculate_dimension() if sheet.max_row else "",
               "rows": sheet.max_row, "columns": sheet.max_column} for sheet in book.worksheets]
    book.close()
    return {"sheets": sheets}


def xlsx_read(source: Path, arguments: dict[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    book = load_workbook(str(source), read_only=True, data_only=True)
    name = arguments.get("sheet") or book.worksheets[0].title
    if name not in book.sheetnames:
        raise ValueError(f"no sheet {name!r}; sheets: {', '.join(book.sheetnames)}")
    sheet = book[name]
    area = (arguments.get("range") or "").replace("$", "").upper()
    if area and not _A1_RE.match(area):
        raise ValueError(f"{area!r} is not an A1 range like B2:F40")
    rows = sheet.iter_rows(**_bounds(area), values_only=False) if area else sheet.iter_rows(values_only=False)
    blocks = []
    for row in rows:
        cells = {f"{get_column_letter(cell.column)}{cell.row}": _value(cell.value) for cell in row
                 if cell.value is not None and hasattr(cell, "column")}
        if cells:
            first = next(iter(cells))
            blocks.append({"ref": f"{name}!{first}", "row": int(re.sub(r"[A-Z]", "", first)), "cells": cells})
    book.close()
    return {"sheet": name, "range": area or "used range", "blocks": blocks}


def _bounds(area: str) -> dict[str, int]:
    from openpyxl.utils import column_index_from_string

    match = _A1_RE.match(area)
    first_col, first_row, last_col, last_row = match.groups()
    return {"min_col": column_index_from_string(first_col), "min_row": int(first_row),
            "max_col": column_index_from_string(last_col or first_col), "max_row": int(last_row or first_row)}


def _value(value: Any) -> Any:
    if isinstance(value, (int, float, str, bool)) or value is None:
        return value
    return str(value)


def csv_info(source: Path) -> dict[str, Any]:
    text = source.read_text(encoding="utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text), dialect=_dialect(text)))
    return {"rows": len(rows), "columns": max((len(row) for row in rows), default=0),
            "header": rows[0] if rows else []}


def csv_read(source: Path, arguments: dict[str, Any]) -> dict[str, Any]:
    text = source.read_text(encoding="utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text), dialect=_dialect(text)))
    numbers = page_numbers(str(arguments.get("range") or ""), len(rows))
    return {"rows_total": len(rows), "header": rows[0] if rows else [],
            "blocks": [{"ref": f"row {number}", "row": number, "values": rows[number - 1]} for number in numbers]}


def _dialect(text: str):
    try:
        return csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        return csv.excel


# --- docx --------------------------------------------------------------------


def _paragraphs(source: Path) -> list[str]:
    with zipfile.ZipFile(source) as archive:
        if archive.getinfo("word/document.xml").file_size > 50 * 1024 * 1024:
            raise ValueError("the document body is too large")
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in root.iter(f"{_W}p"):
        paragraphs.append("".join(node.text or "" for node in paragraph.iter(f"{_W}t")))
    return paragraphs


def docx_info(source: Path) -> dict[str, Any]:
    paragraphs = _paragraphs(source)
    return {"paragraphs": len(paragraphs), "non_empty": sum(1 for text in paragraphs if text.strip())}


def docx_read(source: Path, arguments: dict[str, Any]) -> dict[str, Any]:
    paragraphs = _paragraphs(source)
    numbers = page_numbers(str(arguments.get("range") or ""), len(paragraphs))
    return {"paragraphs_total": len(paragraphs),
            "blocks": [{"ref": f"paragraph {number}", "paragraph": number, "text": paragraphs[number - 1]}
                       for number in numbers if paragraphs[number - 1].strip()]}


INFO = {"pdf": pdf_info, "xlsx": xlsx_info, "csv": csv_info, "docx": docx_info}
READ = {"pdf": pdf_read, "xlsx": xlsx_read, "csv": csv_read, "docx": docx_read}
